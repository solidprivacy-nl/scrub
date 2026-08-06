from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from docx import Document

from document_tools import (
    anonymized_docx_from_original,
    apply_replacements_to_text,
    docx_from_text,
    extract_docx_text,
    replacement_report_csv,
    scrub_report_txt,
)
from processed_text_selection_integration import (
    MANUAL_ROWS_KEY,
    handle_selection_component_event,
    selection_inspection_result,
)
from scrub_key_binding import validate_bound_scrub_key
from scrub_key_bound_export import build_bound_scrub_key
from scrub_key_document_reinsert import reinsert_docx_bytes, reinsert_txt_bytes
from selection_mask_action import processed_text_hash, python_index_to_utf16_offset


SCOPE = "0123456789abcdef"
BINDING = "BABCDEFGHIJKLMNOP"
ORIGINAL_VALUE = "Stichting Zorgpunt"
SOURCE = (
    "Mevrouw Noor werkt bij Stichting Zorgpunt. "
    "Stichting Zorgpunt helpt Noor."
)
TIMESTAMP = datetime(2026, 8, 6, 9, 45, tzinfo=timezone.utc).isoformat().replace(
    "+00:00", "Z"
)


class _UploadedDocx:
    name = "synthetisch_zorgdossier.docx"

    def __init__(self, content: bytes):
        self._content = content

    def getvalue(self) -> bytes:
        return self._content


def _inspect_event() -> dict:
    start = SOURCE.index(ORIGINAL_VALUE)
    end = start + len(ORIGINAL_VALUE)
    return {
        "schema_version": 1,
        "action": "inspect_selection",
        "event_id": "inspect_cross_flow_0001",
        "document_scope_key": SCOPE,
        "processed_text_hash": processed_text_hash(SOURCE),
        "selection": {
            "text": ORIGINAL_VALUE,
            "start_utf16": python_index_to_utf16_offset(SOURCE, start),
            "end_utf16": python_index_to_utf16_offset(SOURCE, end),
            "intersects_marked_content": False,
        },
        "ui_state": {
            "source_scroll_ratio": 0.25,
            "processed_scroll_ratio": 0.25,
        },
    }


def _commit_event(inspection: dict) -> dict:
    return {
        "schema_version": 1,
        "action": "commit_manual_mask",
        "event_id": "commit_cross_flow_0001",
        "inspection_id": inspection["inspection_id"],
        "requested_type": "organization",
        "requested_scope": "all_exact",
        "confirmation_token": inspection["confirmation_token"],
    }


def _selection_row() -> dict:
    session: dict = {}
    inspect_outcome = handle_selection_component_event(
        session,
        _inspect_event(),
        source_text=SOURCE,
        processed_text=SOURCE,
        document_scope_key=SCOPE,
        document_binding_id=BINDING,
        existing_rows=[],
    )
    assert inspect_outcome.status == "ready"

    inspection = selection_inspection_result(session, SCOPE)
    commit_outcome = handle_selection_component_event(
        session,
        _commit_event(inspection),
        source_text=SOURCE,
        processed_text=SOURCE,
        document_scope_key=SCOPE,
        document_binding_id=BINDING,
        existing_rows=[],
    )
    assert commit_outcome.status == "committed"
    assert commit_outcome.row_added is True

    rows = session[MANUAL_ROWS_KEY][SCOPE]
    assert len(rows) == 1
    return dict(rows[0])


def _included_replacements(row: dict) -> dict[str, str]:
    if row.get("include") is not True:
        return {}
    return {str(row["find"]): str(row["replace_with"])}


def _audit_row(row: dict) -> dict:
    return {
        "entity_type": row["entity_type"],
        "detected_text": row["find"],
        "placeholder": row["replace_with"],
        "score": "",
        "source": row["source"],
        "review_status": row["review_status"],
        "review_status_label": row["review_status_label"],
        "reason": row["reason"],
    }


def _bound_key(row: dict) -> dict:
    key_row = {
        "original_value": row["find"],
        "placeholder": row["replace_with"],
        "entity_type": row["entity_type"],
        "type_label": row["type_label"],
        "source": row["source"],
        "review_status": row["review_status"],
        "include": row["include"],
        "timestamp": TIMESTAMP,
    }
    return build_bound_scrub_key([key_row], document_binding_id=BINDING)


def test_selection_row_is_the_same_authoritative_input_for_export_key_and_audit() -> None:
    row = _selection_row()

    assert row["find"] == ORIGINAL_VALUE
    assert row["entity_type"] == "ORGANIZATION"
    assert row["source"] == "manual_selection"
    assert row["source_label"] == "Handmatig uit tekst"
    assert row["selection_scope"] == "all_exact"
    assert row["selection_occurrence_count"] == 2
    assert row["include"] is True

    replacements = _included_replacements(row)
    scrubbed_text = apply_replacements_to_text(SOURCE, replacements)
    placeholder = row["replace_with"]
    assert ORIGINAL_VALUE not in scrubbed_text
    assert scrubbed_text.count(placeholder) == 2

    audit_row = _audit_row(row)
    csv_text = replacement_report_csv([audit_row]).decode("utf-8")
    assert "manual_selection" in csv_text
    assert ORIGINAL_VALUE in csv_text
    assert placeholder in csv_text
    assert "ORGANIZATION" in csv_text

    report_text = scrub_report_txt(
        [audit_row],
        profile="Zorg",
        source_filename="synthetisch_zorgdossier.docx",
    ).decode("utf-8")
    assert "Recognition profile: Zorg" in report_text
    assert "Source file: synthetisch_zorgdossier.docx" in report_text
    assert "- ORGANIZATION: 1" in report_text

    scrub_key = _bound_key(row)
    validation = validate_bound_scrub_key(scrub_key)
    assert validation["ok"] is True, validation
    assert scrub_key["item_count"] == 1
    assert scrub_key["document_binding_id"] == BINDING
    assert scrub_key["items"][0]["original_value"] == ORIGINAL_VALUE
    assert scrub_key["items"][0]["placeholder"] == placeholder
    assert scrub_key["items"][0]["source"] == "manual_selection"


def test_selection_row_roundtrips_through_bound_txt_export_and_reinsert() -> None:
    row = _selection_row()
    scrubbed_text = apply_replacements_to_text(SOURCE, _included_replacements(row))
    scrub_key = _bound_key(row)

    result = reinsert_txt_bytes(scrubbed_text.encode("utf-8"), scrub_key)

    assert result["text"] == SOURCE
    assert result["content_bytes"] == SOURCE.encode("utf-8")
    assert result["replacement_count"] == 2
    assert result["reinserted"] is True
    assert result["replacement_allowed"] is True
    assert result["verified_document_match"] is True
    assert result["legacy_unbound"] is False
    assert result["validation_issues"] == []
    assert result["unknown_placeholders"] == []
    assert result["local_only"] is True
    assert result["ai_processing"] is False
    assert result["cloud_processing"] is False


def test_selection_row_roundtrips_through_original_docx_export_and_reinsert() -> None:
    row = _selection_row()
    replacements = _included_replacements(row)
    scrub_key = _bound_key(row)

    original_docx = docx_from_text(SOURCE)
    scrubbed_docx = anonymized_docx_from_original(
        _UploadedDocx(original_docx),
        replacements,
    )
    scrubbed_text = extract_docx_text(Document(BytesIO(scrubbed_docx)))
    assert ORIGINAL_VALUE not in scrubbed_text
    assert scrubbed_text.count(row["replace_with"]) == 2

    result = reinsert_docx_bytes(scrubbed_docx, scrub_key)
    restored_text = extract_docx_text(Document(BytesIO(result["docx_bytes"])))

    assert restored_text == SOURCE
    assert result["replacement_count"] == 2
    assert result["reinserted"] is True
    assert result["replacement_allowed"] is True
    assert result["verified_document_match"] is True
    assert result["validation_issues"] == []
    assert result["document_type"] == "docx"
    assert result["local_only"] is True
    assert result["ai_processing"] is False
    assert result["cloud_processing"] is False


def test_review_table_include_state_remains_authoritative_for_selection_rows() -> None:
    row = _selection_row()
    row["include"] = False

    assert _included_replacements(row) == {}
    assert apply_replacements_to_text(SOURCE, _included_replacements(row)) == SOURCE

    scrub_key = _bound_key(row)
    validation = validate_bound_scrub_key(scrub_key)
    assert validation["ok"] is True, validation
    assert scrub_key["item_count"] == 0
    assert scrub_key["items"] == []


def test_custom_text_edit_stays_in_document_export_but_fails_closed_for_bound_key() -> None:
    row = _selection_row()
    row["replace_with"] = "Synthetische zorgorganisatie"

    scrubbed_text = apply_replacements_to_text(SOURCE, _included_replacements(row))
    assert scrubbed_text.count("Synthetische zorgorganisatie") == 2
    assert ORIGINAL_VALUE not in scrubbed_text

    scrub_key = _bound_key(row)
    validation = validate_bound_scrub_key(scrub_key)
    assert validation["ok"] is False
    assert "invalid_bound_key" in validation["error_codes"]
